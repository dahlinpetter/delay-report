"""Per day and direction: was ANY train in the window 20+ min late? Refund based on the most delayed one.

    python tools/build_claims.py out.xlsx [start] [end]
"""
import json, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
sys.path.insert(0, __file__.rsplit("/", 1)[0])
from build_excel import journeys, m, hm, header, F, NORM, BOLD, BLUE, GREEN, TITLE, YELLOW

HOME, WORK = "Malmö C", "København H"
LEGS = [("To work", HOME, WORK, ("07:30", "08:30")), ("Home", WORK, HOME, ("15:30", "17:30"))]


def load():
    docs = [json.load(open("data/history.json", encoding="utf-8"))]
    try:
        docs.append(json.load(open("data/delays.json", encoding="utf-8")))
    except FileNotFoundError:
        pass
    by, stops = {}, docs[0]["stops"]
    seen_days = set()
    for doc in docs:
        days = {t["d"] for t in doc["trains"]}
        for t in doc["trains"]:
            if t["d"] in seen_days:
                continue  # history wins for days it covers
            remap = [[stops.index(doc["stops"][s[0]])] + s[1:] for s in t["s"] if doc["stops"][s[0]] in stops]
            by.setdefault(t["d"], []).append({**t, "s": remap})
        seen_days |= days
    return by, stops


def main(out, start, end):
    by, stops = load()
    rows = []
    d = start
    while d <= end:
        if d.isoweekday() <= 5:
            for leg, a, b, w in LEGS:
                js = [j for j in journeys(by, stops, d.isoformat(), a, b) if m(w[0]) <= j["dep"] <= m(w[1])]
                worst = max(js, key=lambda j: j["eff"]) if js else None
                rows.append((d, leg, a, b, len(js), worst))
        d += timedelta(days=1)

    wb = Workbook()
    A = wb.active; A.title = "Assumptions"
    A["A1"] = "Assumptions"; A["A1"].font = TITLE
    header(A, 3, ["Item", "Value", "Note"], [32, 26, 70])
    items = [("30-day ticket price (kr)", 2000, "Assumed, change to your real price"),
             ("Trips per 30-day ticket", 20, "Skånetrafiken: one trip = 1/20 of a 30-day ticket (Region Skåne decision, 2023)"),
             ("Value of one trip (kr)", "=B4/B5", ""),
             ("Level 1 from (min late)", 20, "50% back"), ("Level 1 refund", 0.5, ""),
             ("Level 2 from (min late)", 40, "75% back"), ("Level 2 refund", 0.75, ""),
             ("Level 3 from (min late)", 60, "100% back"), ("Level 3 refund", 1.0, ""),
             ("Morning window", "07:30–08:30 from Malmö C", "Most delayed train departing in the window"),
             ("Evening window", "15:30–17:30 from København H", "Most delayed train departing in the window"),
             ("Period", f"{start} – {end}, weekdays", "Claims can be filed up to about 2 months back"),
             ("Data", "Trafiklab.se KoDa history + live collector, CC BY 4.0", "Delay = arrival delay at your destination; cancelled = wait for next train")]
    for i, (k, v, n) in enumerate(items, start=4):
        A.cell(row=i, column=1, value=k).font = NORM
        c = A.cell(row=i, column=2, value=v); c.font = NORM if isinstance(v, str) and v.startswith("=") else BLUE
        A.cell(row=i, column=3, value=n).font = NORM
    A["B4"].fill = YELLOW; A["B6"].number_format = "#,##0"
    for r in (8, 10, 12):
        A.cell(row=r, column=2).number_format = "0%"
    TV, T1, P1, T2, P2, T3, P3 = ("Assumptions!$B$6", "Assumptions!$B$7", "Assumptions!$B$8",
                                  "Assumptions!$B$9", "Assumptions!$B$10", "Assumptions!$B$11", "Assumptions!$B$12")

    D = wb.create_sheet("Days")
    header(D, 1, ["Date", "Weekday", "Direction", "From", "To", "Trains in window", "Most delayed: departure",
                  "Planned arrival", "Actual arrival", "Minutes late", "Cancelled", "Can claim?", "Refund %", "Refund (kr)"],
           [12, 9, 10, 13, 14, 10, 12, 11, 11, 10, 10, 10, 10, 11])
    red = PatternFill("solid", fgColor="FBE3D6")
    for i, (d, leg, a, b, n, w) in enumerate(rows, start=2):
        vals = [d, d.strftime("%a"), leg, a, b, n]
        vals += ([hm(w["dep"]), hm(w["arr"]), hm(w["arr"] + w["eff"]), int(w["eff"]), "Yes" if w["cancelled"] else ""]
                 if w else ["no data", "", "", None, ""])
        for j, v in enumerate(vals, 1):
            D.cell(row=i, column=j, value=v).font = NORM
        D.cell(row=i, column=1).number_format = "yyyy-mm-dd"
        D.cell(row=i, column=12, value=f'=IF(J{i}="","no data",IF(J{i}>={T1},"Yes","No"))').font = GREEN
        D.cell(row=i, column=13, value=f'=IF(J{i}="",0,IF(J{i}>={T3},{P3},IF(J{i}>={T2},{P2},IF(J{i}>={T1},{P1},0))))').number_format = "0%;;-"
        D.cell(row=i, column=14, value=f"=ROUND(M{i}*{TV},0)").number_format = "#,##0;;-"
        if w and w["eff"] >= 20:
            for j in range(1, 15):
                D.cell(row=i, column=j).fill = red
    last = len(rows) + 1
    D.freeze_panes = "A2"; D.auto_filter.ref = f"A1:N{last}"

    S = wb.create_sheet("Summary", 0)
    S.column_dimensions["A"].width = 46; S.column_dimensions["B"].width = 16
    S["A1"] = "Delay refunds you can still claim"; S["A1"].font = TITLE
    S["A2"] = f"Malmö C ↔ København H, weekdays {start} – {end}. Morning 07:30–08:30, evening 15:30–17:30."; S["A2"].font = NORM
    L, K, N = f"Days!$L$2:$L${last}", f"Days!$C$2:$C${last}", f"Days!$N$2:$N${last}"
    J = f"Days!$J$2:$J${last}"
    lines = [("Total you can claim (kr)", f"=SUM({N})", "#,##0"),
             ("Trips you can claim", f'=COUNTIF({L},"Yes")', "#,##0"),
             ("  to work (mornings)", f'=COUNTIFS({L},"Yes",{K},"To work")', "#,##0"),
             ("  home (evenings)", f'=COUNTIFS({L},"Yes",{K},"Home")', "#,##0"),
             ("  20–39 min late (50%)", f'=COUNTIFS({J},">="&{T1},{J},"<"&{T2})', "#,##0"),
             ("  40–59 min late (75%)", f'=COUNTIFS({J},">="&{T2},{J},"<"&{T3})', "#,##0"),
             ("  60+ min late (100%)", f'=COUNTIFS({J},">="&{T3})', "#,##0"),
             ("Trips checked (days × directions)", f'=COUNTA(Days!$A$2:$A${last})', "#,##0"),
             ("Trips with no train data", f'=COUNTIF({L},"no data")', "#,##0")]
    for i, (k, f, nf) in enumerate(lines, start=4):
        S.cell(row=i, column=1, value=k).font = BOLD if i == 4 else NORM
        c = S.cell(row=i, column=2, value=f); c.number_format = nf
        c.font = Font(name=F, bold=i == 4, size=14 if i == 4 else 10)
    S["A14"] = "Each row on the Days tab is one direction on one day: did any train in your window reach your stop 20+ minutes late? The refund uses the most delayed one."
    S["A14"].font = Font(name=F, italic=True, color="587080")
    wb.save(out)
    print(len(rows), "rows;", sum(1 for r in rows if r[5] and r[5]["eff"] >= 20), "claimable;", sum(1 for r in rows if not r[5]), "no data")


if __name__ == "__main__":
    s = date.fromisoformat(sys.argv[2]) if len(sys.argv) > 2 else date.today() - timedelta(days=61)
    e = date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else date.today()
    main(sys.argv[1], s, e)
