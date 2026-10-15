"""Build the commuter refund analysis workbook from history.json (or delays.json).

    python tools/build_excel.py data/history.json out.xlsx
"""
import json, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

HOME, WORK = "Malmö C", "København H"
OUT_WIN, RET_WIN = ("07:30", "08:30"), ("16:30", "17:30")
PRICE, FACTOR = 2000, 40

F = "Arial"
BOLD = Font(name=F, bold=True)
NORM = Font(name=F)
BLUE = Font(name=F, color="0000FF")
GREEN = Font(name=F, color="008000")
HEAD = Font(name=F, bold=True, color="FFFFFF")
HEADFILL = PatternFill("solid", fgColor="0B3550")
YELLOW = PatternFill("solid", fgColor="FFFF00")
TITLE = Font(name=F, bold=True, size=14)
thin = Side(style="thin", color="D6E0E7")


def m(t):
    h, mm = t.split(":"); return int(h) * 60 + int(mm)


def hm(x):
    x = int(round(x)); return f"{(x // 60) % 24:02d}:{x % 60:02d}"


def journeys(by_day, stops, day, a, b):
    fi, ti = stops.index(a), stops.index(b)
    out = []
    for t in by_day.get(day, []):
        ia = next((i for i, s in enumerate(t["s"]) if s[0] == fi), -1)
        ib = next((i for i, s in enumerate(t["s"]) if s[0] == ti), -1)
        if ia < 0 or ib < 0 or ib <= ia:
            continue
        sb = t["s"][ib]
        out.append({"train": t["n"], "dep": t["s"][ia][1], "arr": sb[3] if len(sb) > 3 else sb[1],
                    "delay": sb[2], "cancelled": bool(t["c"])})
    out.sort(key=lambda j: j["dep"])
    for j in out:
        if not j["cancelled"]:
            j["eff"] = j["delay"]; continue
        nx = next((k for k in out if k["dep"] > j["dep"] and not k["cancelled"]), None)
        j["eff"] = max(0, nx["arr"] + nx["delay"] - j["arr"]) if nx else 60
    return out


def select_trips(doc):
    by_day = {}
    for t in doc["trains"]:
        by_day.setdefault(t["d"], []).append(t)
    days = sorted(by_day)
    start, end = date.fromisoformat(days[0]), date.fromisoformat(days[-1])
    rows, missing = [], []
    d = start
    while d <= end:
        if d.isoweekday() <= 5:
            for leg, a, b, w in (("To work", HOME, WORK, OUT_WIN), ("Home", WORK, HOME, RET_WIN)):
                lo, hi = m(w[0]), m(w[1]); mid = (lo + hi) / 2
                win = [j for j in journeys(by_day, doc["stops"], d.isoformat(), a, b) if lo <= j["dep"] <= hi]
                if not win:
                    missing.append((d.isoformat(), leg)); continue
                j = min(win, key=lambda j: abs(j["dep"] - mid))
                rows.append({"date": d, "leg": leg, "from": a, "to": b, **j})
        d += timedelta(days=1)
    return rows, missing, start, end


def header(ws, row, cols, widths=None):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font, cell.fill = HEAD, HEADFILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def main(src, dst):
    doc = json.load(open(src, encoding="utf-8"))
    rows, missing, start, end = select_trips(doc)
    wb = Workbook()

    # ---------- Assumptions ----------
    A = wb.active; A.title = "Assumptions"
    A["A1"] = "Assumptions"; A["A1"].font = TITLE
    items = [
        ("Home station", HOME, "From Petter's profile defaults"),
        ("Work station", WORK, "From Petter's profile defaults"),
        ("Leave for work between", f"{OUT_WIN[0]}–{OUT_WIN[1]}", "Train closest to the middle of the window is counted"),
        ("Leave work between", f"{RET_WIN[0]}–{RET_WIN[1]}", "Train closest to the middle of the window is counted"),
        ("30-day ticket price (kr)", PRICE, "Assumed, edit to your real price"),
        ("Trips per 30-day ticket", FACTOR, "Assumed divisor, verify with Skånetrafiken"),
        ("Value of one trip (kr)", "=B8/B9", "Formula"),
        ("Level 1: minutes late from", 20, "Skånetrafiken/Öresundståg terms; Lag 2015:953"),
        ("Level 1: refund", 0.5, ""),
        ("Level 2: minutes late from", 40, ""),
        ("Level 2: refund", 0.75, ""),
        ("Level 3: minutes late from", 60, ""),
        ("Level 3: refund", 1.0, ""),
        ("Period analysed", f"{start.isoformat()} – {end.isoformat()} (weekdays)", "Trafiklab KoDa history archive"),
        ("Data source", "Trafiklab.se KoDa (GTFS-Realtime history), CC BY 4.0", "Delay = last live estimate of arrival at your destination"),
        ("Cancelled trains", "Counted as the wait until the next train arrived", ""),
    ]
    header(A, 3, ["Item", "Value", "Note"], [30, 44, 56])
    for i, (k, v, note) in enumerate(items, start=4):
        A.cell(row=i, column=1, value=k).font = NORM
        c = A.cell(row=i, column=2, value=v)
        c.font = NORM if isinstance(v, str) and v.startswith("=") else BLUE
        A.cell(row=i, column=3, value=note).font = NORM
    for r in (8, 9):
        A.cell(row=r, column=2).fill = YELLOW
    A["B10"].number_format = "#,##0.0"
    for r in (12, 14, 16):
        A.cell(row=r, column=2).number_format = "0%"
    # named refs used below
    TV, T1, P1, T2, P2, T3, P3 = "Assumptions!$B$10", "Assumptions!$B$11", "Assumptions!$B$12", "Assumptions!$B$13", "Assumptions!$B$14", "Assumptions!$B$15", "Assumptions!$B$16"

    # ---------- My trips (every counted trip) ----------
    T = wb.create_sheet("My trips")
    cols = ["Date", "Month", "Weekday", "Direction", "From", "To", "Train", "Planned departure", "Planned arrival",
            "Actual arrival", "Minutes late", "Cancelled", "Refund %", "Refund (kr)"]
    header(T, 1, cols, [12, 10, 10, 11, 13, 14, 9, 12, 12, 12, 10, 10, 10, 11])
    for i, r in enumerate(rows, start=2):
        vals = [r["date"], r["date"].strftime("%Y-%m"), r["date"].strftime("%a"), r["leg"], r["from"], r["to"], str(r["train"]) or "–",
                hm(r["dep"]), hm(r["arr"]), hm(r["arr"] + r["eff"]), int(r["eff"]), "Yes" if r["cancelled"] else ""]
        for j, v in enumerate(vals, 1):
            c = T.cell(row=i, column=j, value=v); c.font = NORM
        T.cell(row=i, column=1).number_format = "yyyy-mm-dd"
        T.cell(row=i, column=13, value=f"=IF(K{i}>={T3},{P3},IF(K{i}>={T2},{P2},IF(K{i}>={T1},{P1},0)))").number_format = "0%;;-"
        T.cell(row=i, column=14, value=f"=ROUND(M{i}*{TV},0)").number_format = "#,##0;;-"
        for j in (13, 14):
            T.cell(row=i, column=j).font = GREEN
    last = len(rows) + 1
    T.freeze_panes = "A2"; T.auto_filter.ref = f"A1:N{last}"

    # ---------- Delayed trips ----------
    D = wb.create_sheet("Delayed trips")
    header(D, 1, cols, [12, 10, 10, 11, 13, 14, 9, 12, 12, 12, 10, 10, 10, 11])
    delayed = [(i, r) for i, r in enumerate(rows, start=2) if r["eff"] >= 20]
    for k, (src_row, r) in enumerate(delayed, start=2):
        for j in range(1, 15):
            col = get_column_letter(j)
            ref = f"'My trips'!{col}{src_row}"
            c = D.cell(row=k, column=j, value=f'=IF({ref}="","",{ref})' if j == 12 else f"={ref}")
            c.font = GREEN
        D.cell(row=k, column=1).number_format = "yyyy-mm-dd"
        D.cell(row=k, column=13).number_format = "0%"
        D.cell(row=k, column=14).number_format = "#,##0"
    D.freeze_panes = "A2"
    if delayed:
        D.auto_filter.ref = f"A1:N{len(delayed) + 1}"

    # ---------- Per month ----------
    M = wb.create_sheet("Per month")
    header(M, 1, ["Month", "Trips counted", "20–39 min late", "40–59 min late", "60+ min late", "Delayed trips", "Refund (kr)"], [12, 14, 15, 15, 14, 14, 13])
    months = sorted({r["date"].strftime("%Y-%m") for r in rows})
    R = f"'My trips'!$B$2:$B${last}"; K = f"'My trips'!$K$2:$K${last}"; N = f"'My trips'!$N$2:$N${last}"
    for i, mo in enumerate(months, start=2):
        M.cell(row=i, column=1, value=mo).font = NORM
        M.cell(row=i, column=2, value=f'=COUNTIFS({R},A{i})')
        M.cell(row=i, column=3, value=f'=COUNTIFS({R},A{i},{K},">="&{T1},{K},"<"&{T2})')
        M.cell(row=i, column=4, value=f'=COUNTIFS({R},A{i},{K},">="&{T2},{K},"<"&{T3})')
        M.cell(row=i, column=5, value=f'=COUNTIFS({R},A{i},{K},">="&{T3})')
        M.cell(row=i, column=6, value=f"=SUM(C{i}:E{i})")
        M.cell(row=i, column=7, value=f'=SUMIFS({N},{R},A{i})').number_format = "#,##0"
        for j in range(2, 8):
            M.cell(row=i, column=j).font = NORM
    tr = len(months) + 2
    M.cell(row=tr, column=1, value="Total").font = BOLD
    for j in range(2, 8):
        col = get_column_letter(j)
        c = M.cell(row=tr, column=j, value=f"=SUM({col}2:{col}{tr - 1})"); c.font = BOLD
        c.border = Border(top=Side(style="thin"))
    M.cell(row=tr, column=7).number_format = "#,##0"

    # ---------- Summary ----------
    S = wb.create_sheet("Summary", 0)
    S.column_dimensions["A"].width = 44; S.column_dimensions["B"].width = 18; S.column_dimensions["C"].width = 50
    S["A1"] = "Öresund commute: delay refunds you could have claimed"; S["A1"].font = TITLE
    S["A2"] = f"{HOME} ↔ {WORK}, weekdays {start.isoformat()} – {end.isoformat()}"; S["A2"].font = NORM
    lines = [
        ("Total you could have claimed (kr)", f"=SUM({N})", "#,##0"),
        ("Trips counted", f"=COUNTA('My trips'!$A$2:$A${last})", "#,##0"),
        ("Trips 20+ minutes late", f'=COUNTIFS({K},">="&{T1})', "#,##0"),
        ("Share of trips 20+ minutes late", "=IFERROR(B6/B5,0)", "0.0%"),
        ("  of which 20–39 min (50% back)", f'=COUNTIFS({K},">="&{T1},{K},"<"&{T2})', "#,##0"),
        ("  of which 40–59 min (75% back)", f'=COUNTIFS({K},">="&{T2},{K},"<"&{T3})', "#,##0"),
        ("  of which 60+ min (100% back)", f'=COUNTIFS({K},">="&{T3})', "#,##0"),
        ("Cancelled trains among them", f"=COUNTIFS('My trips'!$L$2:$L${last},\"Yes\",{K},\">=\"&{T1})", "#,##0"),
        ("Average refund per month (kr)", f"=IFERROR(B4/COUNTA('Per month'!$A$2:$A${tr - 1}),0)", "#,##0"),
        ("Value of one trip (kr)", f"={TV}", "#,##0.0"),
    ]
    for i, (k, f, nf) in enumerate(lines, start=4):
        S.cell(row=i, column=1, value=k).font = BOLD if i == 4 else NORM
        c = S.cell(row=i, column=2, value=f); c.number_format = nf; c.font = Font(name=F, bold=(i == 4), size=14 if i == 4 else 10)
    S["C4"] = "Change ticket price and divisor on the Assumptions tab"; S["C4"].font = Font(name=F, italic=True, color="587080")
    # context: every train in the wider rush hours, not just "your" train
    ctx_tot = ctx_late = 0
    by_day = {}
    for t in doc["trains"]:
        by_day.setdefault(t["d"], []).append(t)
    for dd in sorted(by_day):
        if date.fromisoformat(dd).isoweekday() > 5:
            continue
        for a2, b2, w in ((HOME, WORK, ("07:00", "09:00")), (WORK, HOME, ("16:00", "18:00"))):
            for j in journeys(by_day, doc["stops"], dd, a2, b2):
                if m(w[0]) <= j["dep"] <= m(w[1]):
                    ctx_tot += 1; ctx_late += j["eff"] >= 20
    S.cell(row=16, column=1, value="Context: all trains in the rush hours (07–09 and 16–18)").font = BOLD
    S.cell(row=17, column=1, value="Trains checked").font = NORM
    c = S.cell(row=17, column=2, value=ctx_tot); c.font = BLUE; c.number_format = "#,##0"
    S.cell(row=18, column=1, value="Trains 20+ minutes late").font = NORM
    c = S.cell(row=18, column=2, value=ctx_late); c.font = BLUE; c.number_format = "#,##0"
    S.cell(row=19, column=1, value="Share 20+ minutes late").font = NORM
    c = S.cell(row=19, column=2, value="=IFERROR(B18/B17,0)"); c.number_format = "0.0%"; c.font = NORM
    S["C17"] = "Counted from the same Trafiklab history data, both directions"; S["C17"].font = Font(name=F, italic=True, color="587080")
    S.cell(row=21, column=1, value="Data gaps: about 12% of the hourly history files were missing in Trafiklab's archive, so a few delays may be absent.").font = Font(name=F, italic=True, color="587080")
    if missing:
        S.cell(row=15, column=1, value=f"Trips with no train data in the window: {len(missing)} (not counted)").font = Font(name=F, italic=True, color="587080")
    wb.save(dst)
    print(f"{len(rows)} trips, {len(delayed)} delayed, {len(missing)} missing -> {dst}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
